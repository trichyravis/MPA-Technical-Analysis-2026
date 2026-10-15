from dataclasses import asdict
from io import BytesIO
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _safe(d):
    d=d.copy()
    for col in d.select_dtypes(include=["object","string"]).columns:
        d[col]=d[col].map(lambda v:"'"+v if isinstance(v,str) and v.startswith(("=","+","-","@")) else v)
    return d


def excel_bytes(analyzed, settings, metadata, screen=None, swings=None, journal=None):
    """Only pass already cutoff-truncated frames. The export never fetches data."""
    out=BytesIO()
    raw=analyzed[["Open","High","Low","Close","Volume"]].reset_index()
    raw.columns=["Date",*raw.columns[1:]]
    indicators=analyzed.reset_index()
    indicators.columns=["Date",*indicators.columns[1:]]
    notes=pd.DataFrame([
        ["Purpose","Educational daily technical-analysis lab; observations are not recommendations."],
        ["Source",metadata.get("Source","")],
        ["Cutoff",metadata.get("Cutoff","")],
        ["Replay","All price rows, indicators and signals stop at the selected cutoff."],
        ["EMA seed","First full-window SMA; standard alpha=2/(n+1)."],
        ["RSI seed","n changes from n+1 closes; alpha=1/n; flat AG=AL=0 uses 50."],
        ["Support/resistance","Prior-window extrema exclude the tested candle; zones are estimates."],
        ["Signals","Golden/death cross = SMA50/SMA200 event, not a continuing regime."],
        ["Volume","Spot index volume is unavailable and blank; no futures proxy is inserted."],
        ["Replay limits","Current/reference constituents and revised price history are not point-in-time membership/data."],
        ["Data quality","No live feed guarantee. Read the source timestamp and provider terms."],
    ],columns=["Concept","Explanation"])
    sheets={"Read Me":notes,"Parameters":pd.DataFrame(asdict(settings).items(),columns=["Parameter","Value"]),
            "Metadata":pd.DataFrame(metadata.items(),columns=["Field","Value"]),"OHLCV":raw,"Indicators":indicators}
    for name,data in [("Screener",screen),("Confirmed Swings",swings),("Replay Journal",journal)]:
        if data is not None and not data.empty:
            sheets[name]=data.copy()
    with pd.ExcelWriter(out,engine="openpyxl",datetime_format="yyyy-mm-dd",date_format="yyyy-mm-dd") as writer:
        for name,data in sheets.items():
            _safe(data).to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name]
            ws.freeze_panes="A2"
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.fill=PatternFill("solid",fgColor="173D65")
                cell.font=Font(color="FFFFFF",bold=True)
            for column in ws.columns:
                letter=get_column_letter(column[0].column)
                width=max([len(str(c.value)) if c.value is not None else 0 for c in list(column)[:100]])
                ws.column_dimensions[letter].width=min(max(width+2,14),65)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment=Alignment(vertical="top",wrap_text=name in {"Read Me","Metadata"})
                    if isinstance(cell.value,float):
                        cell.number_format="#,##0.0000"
    return out.getvalue()
